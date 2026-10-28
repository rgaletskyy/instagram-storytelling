"""Reads the spreadsheets in src/resources/indices, embeds each row, stores it in Firestore.

Products and content only: `instagram_messages` is written by
routines/messaging_sync from the Graph API, in the shape `message_documents`
builds here.

    uv run rag-index                              # both collections
    uv run rag-index --collection products --limit 20
    uv run rag-index --force                      # re-embed rows already stored

Runs are resumable: a document whose id is already in its collection is
skipped, and rows are written a chunk at a time, so an interrupted run keeps
what it finished. Ids are stable across runs -- the SKU, the Drive file id --
so a rerun never duplicates a row.
"""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import html
import logging
import re
import sys
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from datetime import UTC, date, datetime
from html.parser import HTMLParser
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from domain.contracts import Content, InstagramMessage, Product
from infrastructure.embedding_client import EmbeddingClient
from infrastructure.firestore_client import (
    CONTENT,
    EMBEDDING_TEXT_FIELD,
    PRODUCTS,
    VECTOR_FIELD,
    FirestoreClient,
)

logger = logging.getLogger(__name__)

INDICES_DIR = Path(__file__).resolve().parents[2] / "src" / "resources" / "indices"
# Rows embedded and written together. Small enough that an interruption loses
# little, large enough that a few embedding requests run side by side.
CHUNK = 200
# Seconds to wait between one chunk and the next.
CHUNK_DELAY = 10.0

# What the images index writes in a field that does not apply to the file.
_NOT_APPLICABLE = "—"
# Instagram's own notices and placeholders for a photo, share or reaction.
# They carry no words to search on, and there are thousands of each: embedded,
# they would crowd real messages out of every result.
_DM_NOTICES = {
    "Liked a message",
    "You sent a private reply to a comment on your Instagram post.",
}
_DM_PLACEHOLDER = re.compile(r"^\[[^\]]+\]$")
# Messages are grouped by the day they were sent here, not in UTC: 23:00 UTC is
# already the next morning in Kyiv, and would split one evening's chat in two.
DM_TIMEZONE = ZoneInfo("Europe/Kyiv")

# A file shows a dog when a breed was recognised, or failing that when one of
# its tags is exactly "собака". Exactly: "собака" is also inside "догляд за
# собаками", a topic tag on hundreds of files with no dog in frame. And not
# when another tag says there is no animal, which the describer sometimes adds
# beside "собака" on product shots.
_DOG_TAG = "собака"
_NO_DOG_TAGS = {
    "без тварини",
    "без тварин",
    "без собаки",
    "немає собаки",
    "собака відсутня",
    "відсутній собака",
}


@dataclass(frozen=True)
class IndexSummary:
    indexed: int
    already_stored: int
    no_text: int


@dataclass(frozen=True)
class IndexDocument:
    id: str
    data: dict  # Firestore fields, camelCase, including embeddingText


# --- Reading the spreadsheets --------------------------------------------------


class _TextOnly(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        self.parts.append(data)


def _strip_html(value: str) -> str:
    parser = _TextOnly()
    parser.feed(value)
    return re.sub(r"\s+", " ", html.unescape(" ".join(parser.parts))).strip()


def _text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text == _NOT_APPLICABLE else text


def _number(value) -> float:
    try:
        return float(str(value).replace(",", ".").replace(" ", ""))
    except (TypeError, ValueError):
        return 0.0


def _columns(header: Iterable) -> dict[str, list[int]]:
    """Header name to every column carrying it; the catalogue repeats some."""
    columns: dict[str, list[int]] = {}
    for i, name in enumerate(header):
        if name is not None:
            columns.setdefault(str(name).strip(), []).append(i)
    return columns


def _first(row: tuple, columns: dict[str, list[int]], *names: str, skip=("", "0")) -> str:
    """The first filled cell among the columns named, in the order named."""
    for name in names:
        for i in columns.get(name, []):
            value = _text(row[i]) if i < len(row) else ""
            if value not in skip:
                return value
    return ""


def _embedding_text(*parts: str) -> str:
    return "\n".join(p for p in parts if p)


def _hash_id(*parts: str) -> str:
    return hashlib.sha1("\x1f".join(parts).encode("utf-8")).hexdigest()


def _document(doc_id: str, model, embedding_text: str) -> IndexDocument:
    data = model.model_dump(by_alias=True, exclude={"distance"})
    data[EMBEDDING_TEXT_FIELD] = embedding_text
    return IndexDocument(doc_id, data)


def _first_url(value: str) -> str:
    """The first of a comma-separated list of image links: the main photo."""
    return value.split(",")[0].strip()


# The SKU column under each catalogue export's own name: the old site export
# (Russian headers) and the current B2B export (Ukrainian headers). Every other
# field below likewise names the old column first and the new one after it.
_SKU_COLUMNS = ("Артикул", "SKU")


def read_products(path: Path) -> list[IndexDocument]:
    rows = load_workbook(path, read_only=True).active.iter_rows(values_only=True)
    columns = _columns(next(rows))
    if not any(name in columns for name in _SKU_COLUMNS):
        # Otherwise every row reads as having no SKU and the run quietly
        # indexes nothing.
        raise ValueError(
            f"{path.name} has no SKU column ({' or '.join(_SKU_COLUMNS)}); "
            f"its headers are {sorted(columns)}"
        )
    docs = []
    for row in rows:
        sku = _first(row, columns, *_SKU_COLUMNS)
        if not sku:
            continue
        product = Product(
            sku=sku,
            name=_first(row, columns, "Название модификации (UA)", "Товар/Послуга"),
            price=_number(_first(row, columns, "Цена", "Ціна")),
            brand=_first(row, columns, "Бренд", "Виробник"),
            category=_first(row, columns, "Раздел", "Категорія"),
            description=_strip_html(_first(row, columns, "Описание товара (UA)", "Опис")),
            volume=_first(row, columns, "Об`єм", "Об'єм"),
            # `Вес` is 0 on every row; the real weight is in `Вага`.
            weight=_first(row, columns, "Вага", "Вес"),
            size=_first(row, columns, "Розмір"),
            image_url=_first_url(_first(row, columns, "Зображення")),
        )
        text = _embedding_text(
            product.sku,
            product.name,
            product.brand,
            product.category,
            product.description,
            product.volume,
            product.size,
            product.weight,
        )
        # A SKU may not contain "/", which Firestore reads as a path separator.
        docs.append(_document(sku.replace("/", "_"), product, text))
    return docs


def message_documents(
    messages: Iterable[tuple[str, str, datetime | None, str]],
) -> list[IndexDocument]:
    """`(thread, sender, sent, text)` messages grouped into `instagram_messages` documents.

    One document per conversation and Kyiv day: both sides of the day's chat,
    one message a line as `sender: text`, in the order sent, so a search finds
    a question together with its answer. A single message says little --
    "так", "а на велику породу?" -- and a day of the conversation is what a
    search can actually match. `thread` names the conversation (the customer's
    username) and becomes the document's `sender`; its timestamp is the first
    message's, and its id a hash of conversation and day, so reading the same
    day again gives the same document.
    """
    groups: dict[tuple[str, date | None], list[tuple[datetime | None, str]]] = {}
    for thread, sender, sent, text in messages:
        text = (text or "").strip()
        if not text or text in _DM_NOTICES or _DM_PLACEHOLDER.match(text):
            continue
        day = sent.astimezone(DM_TIMEZONE).date() if sent else None
        # A message's own line breaks would split it over lines of the document.
        line = " ".join(text.split())
        groups.setdefault((thread, day), []).append((sent, f"{sender}: {line}"))

    docs = []
    for (thread, day), lines in groups.items():
        # Stable, so messages sent in the same instant keep the order they came in.
        lines.sort(key=lambda m: m[0] or datetime.max.replace(tzinfo=UTC))
        text = "\n".join(line for _, line in lines)
        message = InstagramMessage(sender=thread, text=text, timestamp=lines[0][0])
        doc_id = _hash_id(thread, day.isoformat() if day else "")
        docs.append(_document(doc_id, message, text))
    return docs


_DRIVE_ID = re.compile(r"(?:/d/|[?&]id=)([\w-]+)")


def _link(cell) -> str:
    """The cell's hyperlink target; the images index shows only "Відкрити"."""
    if cell.hyperlink is not None and cell.hyperlink.target:
        return cell.hyperlink.target
    text = _text(cell.value)
    return text if text.startswith("http") else ""


def _has_dog(breed: str, tags: str) -> bool:
    if breed:
        return True
    tag_set = {tag.strip().lower() for tag in tags.split(",")}
    return _DOG_TAG in tag_set and not tag_set & _NO_DOG_TAGS


def read_content(path: Path) -> list[IndexDocument]:
    # Not read-only: hyperlinks are only loaded in the full mode.
    sheet = load_workbook(path).active
    rows = sheet.iter_rows()
    columns = {name: idx[0] for name, idx in _columns(c.value for c in next(rows)).items()}

    def cell(row, name):
        return row[columns[name]] if name in columns else None

    def value(row, name) -> str:
        c = cell(row, name)
        return _text(c.value) if c is not None else ""

    def link(row, name) -> str:
        c = cell(row, name)
        return _link(c) if c is not None else ""

    docs = []
    for row in rows:
        file_name = value(row, "Файл")
        if not file_name:
            continue
        drive_url = link(row, "Google Drive")
        content = Content(
            file_name=file_name,
            drive_url=drive_url,
            description=value(row, "Опис"),
            screenshots_path=link(row, "Screenshots"),
            audio_transcribe=value(row, "AudioTranscribe"),
            product=value(row, "Продукт"),
            brand=value(row, "Бренд"),
            breed=value(row, "Порода"),
            tags=value(row, "Теги"),
            type=value(row, "Тип"),
            text_on_image=value(row, "Текст на фото"),
            size_kb=_number(value(row, "Розмір, KB")),
            has_dog=_has_dog(value(row, "Порода"), value(row, "Теги")),
        )
        text = _embedding_text(
            content.description,
            content.audio_transcribe,
            content.product,
            content.brand,
            content.breed,
            content.tags,
            content.type,
            content.text_on_image,
        )
        match = _DRIVE_ID.search(drive_url)
        doc_id = match.group(1) if match else _hash_id(value(row, "Папка"), file_name)
        docs.append(_document(doc_id, content, text))
    return docs


SOURCES: dict[str, tuple[str, Callable[[Path], list[IndexDocument]]]] = {
    PRODUCTS: ("products.xlsx", read_products),
    CONTENT: ("images_index.xlsx", read_content),
}


# --- Embedding and storing -----------------------------------------------------


async def index_collection(
    collection: str,
    docs: list[IndexDocument],
    firestore: FirestoreClient,
    embedder: EmbeddingClient,
    *,
    limit: int | None = None,
    force: bool = False,
) -> IndexSummary:
    """Embed and store the documents not already in `collection`."""
    # The same row twice in one sheet would be embedded twice and written once.
    unique = list({d.id: d for d in docs}.values())
    # A row nobody has described yet -- the images index has a few hundred --
    # has nothing to search on, and an empty input fails the whole request.
    described = [d for d in unique if d.data[EMBEDDING_TEXT_FIELD]]
    stored = set() if force else await firestore.existing_ids(collection)
    pending = [d for d in described if d.id not in stored]
    if limit is not None:
        pending = pending[:limit]

    indexed = 0
    for start in range(0, len(pending), CHUNK):
        chunk = pending[start : start + CHUNK]
        vectors = await embedder.embed_documents([d.data[EMBEDDING_TEXT_FIELD] for d in chunk])
        await firestore.upsert(
            collection,
            ((d.id, {**d.data, VECTOR_FIELD: v}) for d, v in zip(chunk, vectors, strict=True)),
        )
        indexed += len(chunk)
        logger.info("%s: %d/%d indexed", collection, indexed, len(pending))
        if start + CHUNK < len(pending):
            await asyncio.sleep(CHUNK_DELAY)
    return IndexSummary(
        indexed=indexed,
        already_stored=sum(1 for d in described if d.id in stored),
        no_text=len(unique) - len(described),
    )


async def run(
    collections: list[str],
    indices_dir: Path = INDICES_DIR,
    *,
    limit: int | None = None,
    force: bool = False,
    firestore: FirestoreClient | None = None,
    embedder: EmbeddingClient | None = None,
) -> dict[str, IndexSummary]:
    firestore = firestore or FirestoreClient()
    embedder = embedder or EmbeddingClient()
    summary = {}
    for collection in collections:
        file_name, read = SOURCES[collection]
        docs = read(indices_dir / file_name)
        summary[collection] = await index_collection(
            collection, docs, firestore, embedder, limit=limit, force=force
        )
    return summary


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="rag-index", description=__doc__.splitlines()[0])
    parser.add_argument(
        "--collection",
        action="append",
        choices=list(SOURCES),
        help="index only this collection; repeatable (default: both)",
    )
    parser.add_argument("--limit", type=int, help="index at most N new rows per collection")
    parser.add_argument("--force", action="store_true", help="re-embed rows already stored")
    parser.add_argument("--dir", type=Path, default=INDICES_DIR, help="folder of the spreadsheets")
    args = parser.parse_args(argv)

    from infrastructure.env import load_dotenv

    load_dotenv()
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    summary = asyncio.run(
        run(args.collection or list(SOURCES), args.dir, limit=args.limit, force=args.force)
    )
    for collection, s in summary.items():
        print(
            f"{collection}: {s.indexed} indexed, {s.already_stored} already stored, "
            f"{s.no_text} with no text to embed"
        )
    return 0


if __name__ == "__main__":
    sys.exit(main())
