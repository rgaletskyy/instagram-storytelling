"""Reads the spreadsheets in src/resources/indices, embeds each row, stores it in Firestore.

    uv run rag-index                              # all three collections
    uv run rag-index --collection products --limit 20
    uv run rag-index --force                      # re-embed rows already stored

Runs are resumable: a document whose id is already in its collection is
skipped, and rows are written a chunk at a time, so an interrupted run over the
messages keeps what it finished. Ids are stable across runs -- the SKU, the
Drive file id, a hash of conversation, sender and day -- so a rerun never
duplicates a row.
"""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import html
import logging
import re
import sys
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from datetime import UTC, date, datetime
from html.parser import HTMLParser
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from domain.contracts import Content, InstagramMessage, Product
from infrastructure.embedding_client import EmbeddingClient
from infrastructure.firestore_client import (
    CONTENT,
    EMBEDDING_TEXT_FIELD,
    INSTAGRAM_MESSAGES,
    PRODUCTS,
    VECTOR_FIELD,
    FirestoreClient,
)

logger = logging.getLogger(__name__)

INDICES_DIR = Path(__file__).resolve().parents[1] / "resources" / "indices"
# Rows embedded and written together. Small enough that an interruption loses
# little, large enough that a few embedding requests run side by side.
CHUNK = 200
# Seconds to wait between one chunk and the next.
CHUNK_DELAY = 10.0

# What the images index writes in a field that does not apply to the file.
_NOT_APPLICABLE = "—"
# Instagram's own notices and placeholders for a photo, share or reaction.
# They carry no words to search on, and there are thousands of each: embedded,
# they would crowd real messages out of every result.
_DM_NOTICES = {
    "Liked a message",
    "You sent a private reply to a comment on your Instagram post.",
}
_DM_PLACEHOLDER = re.compile(r"^\[[^\]]+\]$")
# Messages are grouped by the day they were sent here, not in UTC: 23:00 UTC is
# already the next morning in Kyiv, and would split one evening's chat in two.
_DM_TIMEZONE = ZoneInfo("Europe/Kyiv")


@dataclass(frozen=True)
class IndexSummary:
    indexed: int
    already_stored: int
    no_text: int


@dataclass(frozen=True)
class IndexDocument:
    id: str
    data: dict  # Firestore fields, camelCase, including embeddingText


# --- Reading the spreadsheets --------------------------------------------------


class _TextOnly(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        self.parts.append(data)


def _strip_html(value: str) -> str:
    parser = _TextOnly()
    parser.feed(value)
    return re.sub(r"\s+", " ", html.unescape(" ".join(parser.parts))).strip()


def _text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text == _NOT_APPLICABLE else text


def _number(value) -> float:
    try:
        return float(str(value).replace(",", ".").replace(" ", ""))
    except (TypeError, ValueError):
        return 0.0


def _columns(header: Iterable) -> dict[str, list[int]]:
    """Header name to every column carrying it; the catalogue repeats some."""
    columns: dict[str, list[int]] = {}
    for i, name in enumerate(header):
        if name is not None:
            columns.setdefault(str(name).strip(), []).append(i)
    return columns


def _first(row: tuple, columns: dict[str, list[int]], *names: str, skip=("", "0")) -> str:
    """The first filled cell among the columns named, in the order named."""
    for name in names:
        for i in columns.get(name, []):
            value = _text(row[i]) if i < len(row) else ""
            if value not in skip:
                return value
    return ""


def _embedding_text(*parts: str) -> str:
    return "\n".join(p for p in parts if p)


def _hash_id(*parts: str) -> str:
    return hashlib.sha1("\x1f".join(parts).encode("utf-8")).hexdigest()


def _document(doc_id: str, model, embedding_text: str) -> IndexDocument:
    data = model.model_dump(by_alias=True, exclude={"distance"})
    data[EMBEDDING_TEXT_FIELD] = embedding_text
    return IndexDocument(doc_id, data)


def read_products(path: Path) -> list[IndexDocument]:
    rows = load_workbook(path, read_only=True).active.iter_rows(values_only=True)
    columns = _columns(next(rows))
    docs = []
    for row in rows:
        sku = _first(row, columns, "Артикул")
        if not sku:
            continue
        product = Product(
            sku=sku,
            name=_first(row, columns, "Название (UA)", "Название модификации (UA)"),
            price=_number(_first(row, columns, "Цена")),
            brand=_first(row, columns, "Бренд"),
            category=_first(row, columns, "Раздел"),
            description=_strip_html(_first(row, columns, "Описание товара (UA)")),
            volume=_first(row, columns, "Об`єм", "Об'єм"),
            # `Вес` is 0 on every row; the real weight is in `Вага`.
            weight=_first(row, columns, "Вага", "Вес"),
            size=_first(row, columns, "Розмір"),
        )
        text = _embedding_text(
            product.sku,
            product.name,
            product.brand,
            product.category,
            product.description,
            product.volume,
            product.size,
            product.weight,
        )
        # A SKU may not contain "/", which Firestore reads as a path separator.
        docs.append(_document(sku.replace("/", "_"), product, text))
    return docs


def read_instagram_messages(path: Path) -> list[IndexDocument]:
    """One document per conversation, sender and day, its messages in the order sent.

    A single message says little -- "так", "а на велику породу?" -- and a day of
    them from one person is what a search can actually match. Grouping by
    conversation as well as sender keeps the shop's own replies, which it sends
    to everyone, from merging across customers. The document's timestamp is its
    first message's.
    """
    rows = load_workbook(path, read_only=True).active.iter_rows(values_only=True)
    columns = _columns(next(rows))
    groups: dict[tuple[str, str, date | None], list[tuple[datetime | None, str]]] = {}
    for row in rows:
        text = _first(row, columns, "Message")
        if not text or text in _DM_NOTICES or _DM_PLACEHOLDER.match(text):
            continue
        sender = _first(row, columns, "Sender")
        thread = _first(row, columns, "Thread Folder")
        try:
            sent = datetime.fromisoformat(_first(row, columns, "Timestamp UTC")).replace(tzinfo=UTC)
        except ValueError:
            sent = None
        day = sent.astimezone(_DM_TIMEZONE).date() if sent else None
        groups.setdefault((thread, sender, day), []).append((sent, text))

    docs = []
    for (thread, sender, day), messages in groups.items():
        # Stable, so messages sent in the same instant keep the sheet's order.
        messages.sort(key=lambda m: m[0] or datetime.max.replace(tzinfo=UTC))
        text = "\n".join(t for _, t in messages)
        message = InstagramMessage(sender=sender, text=text, timestamp=messages[0][0])
        doc_id = _hash_id(thread, sender, day.isoformat() if day else "")
        docs.append(_document(doc_id, message, text))
    return docs


_DRIVE_ID = re.compile(r"(?:/d/|[?&]id=)([\w-]+)")


def _link(cell) -> str:
    """The cell's hyperlink target; the images index shows only "Відкрити"."""
    if cell.hyperlink is not None and cell.hyperlink.target:
        return cell.hyperlink.target
    text = _text(cell.value)
    return text if text.startswith("http") else ""


def read_content(path: Path) -> list[IndexDocument]:
    # Not read-only: hyperlinks are only loaded in the full mode.
    sheet = load_workbook(path).active
    rows = sheet.iter_rows()
    columns = {name: idx[0] for name, idx in _columns(c.value for c in next(rows)).items()}

    def cell(row, name):
        return row[columns[name]] if name in columns else None

    def value(row, name) -> str:
        c = cell(row, name)
        return _text(c.value) if c is not None else ""

    def link(row, name) -> str:
        c = cell(row, name)
        return _link(c) if c is not None else ""

    docs = []
    for row in rows:
        file_name = value(row, "Файл")
        if not file_name:
            continue
        drive_url = link(row, "Google Drive")
        content = Content(
            file_name=file_name,
            drive_url=drive_url,
            description=value(row, "Опис"),
            screenshots_path=link(row, "Screenshots"),
            audio_transcribe=value(row, "AudioTranscribe"),
            product=value(row, "Продукт"),
            brand=value(row, "Бренд"),
            breed=value(row, "Порода"),
            tags=value(row, "Теги"),
            type=value(row, "Тип"),
            text_on_image=value(row, "Текст на фото"),
            size_kb=_number(value(row, "Розмір, KB")),
        )
        text = _embedding_text(
            content.description,
            content.audio_transcribe,
            content.product,
            content.brand,
            content.breed,
            content.tags,
            content.type,
            content.text_on_image,
        )
        match = _DRIVE_ID.search(drive_url)
        doc_id = match.group(1) if match else _hash_id(value(row, "Папка"), file_name)
        docs.append(_document(doc_id, content, text))
    return docs


SOURCES: dict[str, tuple[str, Callable[[Path], list[IndexDocument]]]] = {
    PRODUCTS: ("products.xlsx", read_products),
    INSTAGRAM_MESSAGES: ("Instagram_DM_Index.xlsx", read_instagram_messages),
    CONTENT: ("images_index.xlsx", read_content),
}


# --- Embedding and storing -----------------------------------------------------


async def index_collection(
    collection: str,
    docs: list[IndexDocument],
    firestore: FirestoreClient,
    embedder: EmbeddingClient,
    *,
    limit: int | None = None,
    force: bool = False,
) -> IndexSummary:
    """Embed and store the documents not already in `collection`."""
    # The same row twice in one sheet would be embedded twice and written once.
    unique = list({d.id: d for d in docs}.values())
    # A row nobody has described yet -- the images index has a few hundred --
    # has nothing to search on, and an empty input fails the whole request.
    described = [d for d in unique if d.data[EMBEDDING_TEXT_FIELD]]
    stored = set() if force else await firestore.existing_ids(collection)
    pending = [d for d in described if d.id not in stored]
    if limit is not None:
        pending = pending[:limit]

    indexed = 0
    for start in range(0, len(pending), CHUNK):
        chunk = pending[start : start + CHUNK]
        vectors = await embedder.embed_documents([d.data[EMBEDDING_TEXT_FIELD] for d in chunk])
        await firestore.upsert(
            collection,
            ((d.id, {**d.data, VECTOR_FIELD: v}) for d, v in zip(chunk, vectors, strict=True)),
        )
        indexed += len(chunk)
        logger.info("%s: %d/%d indexed", collection, indexed, len(pending))
        if start + CHUNK < len(pending):
            await asyncio.sleep(CHUNK_DELAY)
    return IndexSummary(
        indexed=indexed,
        already_stored=sum(1 for d in described if d.id in stored),
        no_text=len(unique) - len(described),
    )


async def run(
    collections: list[str],
    indices_dir: Path = INDICES_DIR,
    *,
    limit: int | None = None,
    force: bool = False,
    firestore: FirestoreClient | None = None,
    embedder: EmbeddingClient | None = None,
) -> dict[str, IndexSummary]:
    firestore = firestore or FirestoreClient()
    embedder = embedder or EmbeddingClient()
    summary = {}
    for collection in collections:
        file_name, read = SOURCES[collection]
        docs = read(indices_dir / file_name)
        summary[collection] = await index_collection(
            collection, docs, firestore, embedder, limit=limit, force=force
        )
    return summary


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="rag-index", description=__doc__.splitlines()[0])
    parser.add_argument(
        "--collection",
        action="append",
        choices=list(SOURCES),
        help="index only this collection; repeatable (default: all three)",
    )
    parser.add_argument("--limit", type=int, help="index at most N new rows per collection")
    parser.add_argument("--force", action="store_true", help="re-embed rows already stored")
    parser.add_argument("--dir", type=Path, default=INDICES_DIR, help="folder of the spreadsheets")
    args = parser.parse_args(argv)

    from instagram_marketing_agent.config import load_dotenv

    load_dotenv()
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    summary = asyncio.run(
        run(args.collection or list(SOURCES), args.dir, limit=args.limit, force=args.force)
    )
    for collection, s in summary.items():
        print(
            f"{collection}: {s.indexed} indexed, {s.already_stored} already stored, "
            f"{s.no_text} with no text to embed"
        )
    return 0


if __name__ == "__main__":
    sys.exit(main())
