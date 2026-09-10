"""Cataloguing the image library: what is read, what is skipped, what is written."""

import asyncio
import io

import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image

from instagram_marketing_agent import drive, image_index
from instagram_marketing_agent.drive import file_id
from instagram_marketing_agent.image_index import (
    COLUMNS,
    EMPTY,
    INDEX_SHEET,
    INVENTORY_SHEET,
    PARALLELISM,
    ImageFacts,
    read_inventory,
)

pytestmark = pytest.mark.unit

DRIVE = "https://drive.google.com/file/d/{}/view?usp=drivesdk"


def _jpeg(path, size=(32, 32)):
    """A real JPEG on disk, small enough to be free."""
    buffer = io.BytesIO()
    Image.new("RGB", size, "red").save(buffer, "JPEG")
    path.write_bytes(buffer.getvalue())
    return path



def _inventory(tmp_path, rows):
    """An inventory in the real one's shape: links live in the cell, not the text."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INVENTORY_SHEET
    sheet.append(
        ["#", "Папка (верхній рівень)", "Шлях", "Файл", "Тип", "Розмір, KB",
         "Дата зміни", "Локальний файл", "Google Drive", "Рекомендовано (пілот)"]
    )
    for number, (folder, name, kind, recommended, drive_id) in enumerate(rows, 1):
        sheet.append(
            [number, folder, f"{folder}/{name}", name, kind, 100.5, "2024-06-11",
             "Відкрити", "Відкрити", recommended]
        )
        sheet.cell(row=number + 1, column=8).hyperlink = f"file:///G:/{folder}/{name}"
        sheet.cell(row=number + 1, column=9).hyperlink = DRIVE.format(drive_id)

    tmp_path.mkdir(parents=True, exist_ok=True)
    path = tmp_path / "index_inventory.xlsx"
    workbook.save(path)
    return path


@pytest.fixture
def indexed(monkeypatch):
    """Stand in for Drive and the vision model; record what each was asked for."""
    seen = {"downloads": [], "described": 0}

    async def fake_download(client, url):
        seen["downloads"].append(url)
        buffer = io.BytesIO()
        Image.new("RGB", (64, 64), "red").save(buffer, "JPEG")
        return buffer.getvalue()

    async def fake_describe(image, model=None):
        seen["described"] += 1
        return ImageFacts(
            description="Рудий пес на дивані",
            breed="такса",
            tags=["собака", "такса", "лайфстайл"],
            kind="лайфстайл",
        )

    monkeypatch.setattr(drive, "download", fake_download)
    monkeypatch.setattr(image_index, "describe", fake_describe)
    return seen


@pytest.mark.parametrize(
    "url,expected",
    [
        ("https://drive.google.com/file/d/1zfNagl-Ut1U/view?usp=drivesdk", "1zfNagl-Ut1U"),
        ("https://drive.google.com/uc?export=download&id=1KDImHUh25", "1KDImHUh25"),
        ("https://drive.google.com/open?id=1sLDqCZ&authuser=0", "1sLDqCZ"),
    ],
)
def test_the_file_id_is_read_out_of_either_link_shape(url, expected):
    assert file_id(url) == expected


def test_a_link_with_no_id_is_refused():
    with pytest.raises(ValueError, match="no Drive file id"):
        file_id("https://example.com/photo.jpg")


def test_only_recommended_files_are_read_by_default(tmp_path):
    inventory = _inventory(
        tmp_path,
        [
            ("Grooming", "a.jpg", "зображення", "так", "id_a"),
            ("Grooming", "b.jpg", "зображення", None, "id_b"),
            ("Grooming", "clip.mp4", "відео", "так", "id_c"),
        ],
    )

    assert [r.file for r in read_inventory(inventory)] == ["a.jpg", "clip.mp4"]
    assert [r.file for r in read_inventory(inventory, only_recommended=False)] == [
        "a.jpg",
        "b.jpg",
        "clip.mp4",
    ]


def test_video_is_read_and_stray_files_are_not(tmp_path):
    """The inventory's own type column, and the extension, both have to agree."""
    inventory = _inventory(
        tmp_path,
        [
            ("B2B", "clip.mp4", "відео", "так", "id_a"),
            ("B2B", "notes.pdf", "зображення", "так", "id_b"),
            ("B2B", "sheet.xlsx", "інше", "так", "id_c"),
            ("B2B", "photo.HEIC", "зображення", "так", "id_d"),
        ],
    )

    rows = read_inventory(inventory)

    assert [(r.file, r.kind) for r in rows] == [("clip.mp4", "video"), ("photo.HEIC", "img")]
    assert rows[0].is_video and not rows[1].is_video


def test_a_folder_can_be_picked_out(tmp_path):
    inventory = _inventory(
        tmp_path,
        [
            ("Dogs_from_Grooming", "a.jpg", "зображення", "так", "id_a"),
            ("B2B-партнери", "b.jpg", "зображення", "так", "id_b"),
        ],
    )

    rows = read_inventory(inventory, folder="grooming")
    assert [r.file for r in rows] == ["a.jpg"]
    assert rows[0].drive_url == DRIVE.format("id_a")


async def test_the_index_is_written_in_the_pilot_shape(tmp_path, indexed):
    inventory = _inventory(
        tmp_path, [("Grooming", "a.jpg", "зображення", "так", "id_a")]
    )
    out = tmp_path / "images_index.xlsx"

    count, failures = await image_index.build_index(inventory, out)

    assert (count, failures) == (1, [])
    sheet = load_workbook(out)[INDEX_SHEET]
    assert [c.value for c in sheet[1]] == list(COLUMNS)
    written = dict(zip(COLUMNS, (c.value for c in sheet[2]), strict=True))
    assert written["Файл"] == "a.jpg"
    assert written["Папка"] == "Grooming"
    assert written["Тип"] == "img"
    assert written["Опис"] == "Рудий пес на дивані"
    # Absent fields carry the pilot's dash rather than an empty cell.
    assert written["Продукт"] == EMPTY and written["Бренд"] == EMPTY
    assert written["Теги"] == "собака, такса, лайфстайл"
    assert written["Категорія"] == "лайфстайл"
    # An image has neither of the video columns.
    assert written["Screenshots"] == EMPTY and written["AudioTranscribe"] == EMPTY
    # The link column stays clickable, as it is in the inventory.
    drive_column = COLUMNS.index("Google Drive") + 1
    assert sheet.cell(row=2, column=drive_column).hyperlink.target == DRIVE.format("id_a")


@pytest.mark.parametrize(
    "seconds,every",
    [(1.0, 5.0), (59.9, 5.0), (60.0, 5.0), (60.1, 8.0), (120.0, 8.0), (120.1, None),
     (600.0, None)],
)
def test_how_often_a_clip_is_sampled(seconds, every):
    """Every 5s under a minute, every 8s to two, and past that not at all."""
    assert image_index.screenshot_interval(seconds) == every


def test_frames_are_compiled_into_one_entry():
    """Kept whole and numbered: a clip is sampled because it changes."""
    merged = image_index.merge_facts(
        [
            ImageFacts(description="пес біжить", tags=["собака", "рух"], kind="лайфстайл"),
            ImageFacts(
                description="банка на столі",
                product="Calming",
                brand="Natural Dog Company",
                tags=["товар", "собака"],
                kind="продуктове",
                text_in_image="Calming",
            ),
            ImageFacts(description="пес спить", breed="мопс", tags=["сон"], kind="лайфстайл"),
        ]
    )

    assert merged.description == (
        "Кадр 1: пес біжить\nКадр 2: банка на столі\nКадр 3: пес спить"
    )
    # First answer wins for a single-valued field; tags are the union, in order.
    assert (merged.product, merged.brand, merged.breed) == (
        "Calming",
        "Natural Dog Company",
        "мопс",
    )
    assert merged.tags == ["собака", "рух", "товар", "сон"]
    assert merged.kind == "лайфстайл"
    assert merged.text_in_image == "Calming"


def test_compiling_nothing_is_not_an_error():
    """A clip whose frames all failed still gets a row, with empty fields."""
    assert image_index.merge_facts([]).description == ""


@pytest.fixture
def clip(monkeypatch, tmp_path):
    """Stand in for ffmpeg, Drive's upload half and the transcriber."""
    seen = {
        "interval": None,
        "folder": None,
        "parent": None,
        "asked_where": None,
        "uploaded": [],
        "described": 0,
    }

    async def fake_frames(video, out_dir, every_seconds):
        seen["interval"] = every_seconds
        out_dir.mkdir(parents=True, exist_ok=True)
        frames = []
        for i in (1, 2):
            frame = out_dir / f"frame_{i:03d}.jpg"
            _jpeg(frame)
            frames.append(frame)
        return frames

    async def fake_audio(video, out_path):
        out_path.write_bytes(b"wav")
        return out_path

    async def fake_transcribe(audio):
        return "привіт, це тест"

    async def fake_parent(client, url_or_id):
        seen["asked_where"] = url_or_id
        return "the-clips-own-folder"

    async def fake_folder(client, name, parent):
        seen["folder"] = name
        seen["parent"] = parent
        return "folder-id", drive.folder_link("folder-id")

    async def fake_upload(client, path, parent_id):
        seen["uploaded"].append((path.name, parent_id))
        return "file-id"

    async def fake_describe(image, model=None):
        seen["described"] += 1
        return ImageFacts(description=f"кадр {image.name}", tags=["собака"])

    monkeypatch.setattr(image_index.ffmpeg, "duration", lambda video: 30.0)
    monkeypatch.setattr(image_index.ffmpeg, "extract_frames_every", fake_frames)
    monkeypatch.setattr(image_index.ffmpeg, "extract_audio", fake_audio)
    monkeypatch.setattr(image_index.llm, "transcribe_audio", fake_transcribe)
    monkeypatch.setattr(drive, "parent_of", fake_parent)
    monkeypatch.setattr(drive, "public_folder", fake_folder)
    monkeypatch.setattr(drive, "upload", fake_upload)
    monkeypatch.setattr(image_index, "describe", fake_describe)
    return seen


async def test_a_clip_is_sampled_uploaded_and_transcribed(tmp_path, clip):
    item = await image_index.index_video(
        None, tmp_path / "IMG_1.MOV", "IMG_1.MOV", "m", DRIVE.format("id_v")
    )

    assert clip["interval"] == 5.0
    # The folder is named for the clip and created beside it, not somewhere
    # configured: the library keeps its own shape.
    assert clip["asked_where"] == DRIVE.format("id_v")
    assert clip["parent"] == "the-clips-own-folder"
    assert clip["folder"] == "IMG_1"
    assert [name for name, _ in clip["uploaded"]] == [
        "frame_001.jpg",
        "frame_002.jpg",
    ]
    assert all(parent == "folder-id" for _, parent in clip["uploaded"])
    assert item.screenshots_url == drive.folder_link("folder-id")
    assert item.transcript == "привіт, це тест"
    assert item.facts.description.startswith("Кадр 1: кадр frame_001.jpg")


async def test_a_long_clip_is_transcribed_but_not_sampled(tmp_path, clip, monkeypatch):
    """Past two minutes the frames run into dozens and say the same thing."""
    monkeypatch.setattr(image_index.ffmpeg, "duration", lambda video: 300.0)

    item = await image_index.index_video(None, tmp_path / "long.mp4", "long.mp4", "m")

    assert clip["interval"] is None
    assert clip["uploaded"] == []
    assert item.screenshots_url == ""
    assert item.facts.description == ""
    # The words are still worth having.
    assert item.transcript == "привіт, це тест"


async def test_a_clip_with_no_audio_still_indexes(tmp_path, clip, monkeypatch):
    async def no_audio(video, out_path):
        return None

    monkeypatch.setattr(image_index.ffmpeg, "extract_audio", no_audio)

    item = await image_index.index_video(None, tmp_path / "silent.mp4", "silent.mp4", "m")

    assert item.transcript == ""
    assert item.facts.description != ""


async def test_a_transcription_failure_does_not_lose_the_frames(tmp_path, clip, monkeypatch):
    async def boom(audio):
        raise RuntimeError("gemini is down")

    monkeypatch.setattr(image_index.llm, "transcribe_audio", boom)

    item = await image_index.index_video(None, tmp_path / "IMG_2.MOV", "IMG_2.MOV", "m")

    assert item.transcript == ""
    assert item.screenshots_url == drive.folder_link("folder-id")


async def test_a_video_row_is_written_with_its_screenshots_and_transcript(
    tmp_path, clip, monkeypatch
):
    inventory = _inventory(
        tmp_path, [("UGC", "IMG_1.MOV", "відео", "так", "id_v")]
    )
    out = tmp_path / "images_index.xlsx"

    async def fake_download(client, url):
        return b"not really a video"

    monkeypatch.setattr(drive, "download", fake_download)

    count, failures = await image_index.build_index(inventory, out)

    assert (count, failures) == (1, [])
    sheet = load_workbook(out)[INDEX_SHEET]
    written = dict(zip(COLUMNS, (c.value for c in sheet[2]), strict=True))
    assert written["Тип"] == "video"
    assert written["AudioTranscribe"] == "привіт, це тест"
    assert written["Screenshots"] == "Відкрити"
    column = COLUMNS.index("Screenshots") + 1
    assert sheet.cell(row=2, column=column).hyperlink.target == drive.folder_link(
        "folder-id"
    )


async def test_a_name_already_in_the_index_is_skipped_whatever_its_folder(
    tmp_path, indexed
):
    """Matching is on 'Файл' alone, so the same name elsewhere counts as done."""
    first = _inventory(
        tmp_path / "one", [("Grooming", "IMG_0830.HEIC", "зображення", "так", "id_a")]
    )
    second = _inventory(
        tmp_path / "two", [("UGC", "IMG_0830.HEIC", "зображення", "так", "id_b")]
    )
    out = tmp_path / "images_index.xlsx"

    assert (await image_index.build_index(first, out))[0] == 1
    assert (await image_index.build_index(second, out))[0] == 0
    assert indexed["described"] == 1


async def test_one_name_twice_in_a_run_is_described_once(tmp_path, indexed):
    """237 filenames repeat in the real inventory, mostly between subfolders."""
    inventory = _inventory(
        tmp_path,
        [
            ("Сток", "1.jpg", "зображення", "так", "id_a"),
            ("Сток", "1.jpg", "зображення", "так", "id_b"),
            ("Сток", "2.jpg", "зображення", "так", "id_c"),
        ],
    )
    out = tmp_path / "images_index.xlsx"

    count, _ = await image_index.build_index(inventory, out)

    assert count == 2
    assert [c.value for c in load_workbook(out)[INDEX_SHEET]["B"]][1:] == [
        "1.jpg",
        "2.jpg",
    ]


def test_the_file_column_is_found_by_its_header_not_its_position(tmp_path):
    """A re-ordered or hand-edited index must not be matched on the wrong column."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INDEX_SHEET
    sheet.append(["Файл", "#", "Папка", "Опис"])
    sheet.append(["a.jpg", 1, "Grooming", "опис"])
    out = tmp_path / "images_index.xlsx"
    workbook.save(out)

    assert image_index.indexed_filenames(out) == {"a.jpg"}


def test_an_index_without_the_file_column_is_refused(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INDEX_SHEET
    sheet.append(["#", "Папка", "Опис"])
    out = tmp_path / "images_index.xlsx"
    workbook.save(out)

    with pytest.raises(ValueError, match="no 'Файл' column"):
        image_index.indexed_filenames(out)


def test_no_index_yet_means_nothing_is_skipped(tmp_path):
    assert image_index.indexed_filenames(tmp_path / "not-there.xlsx") == set()


async def test_a_second_run_only_describes_what_is_missing(tmp_path, indexed):
    inventory = _inventory(
        tmp_path,
        [
            ("Grooming", "a.jpg", "зображення", "так", "id_a"),
            ("Grooming", "b.jpg", "зображення", "так", "id_b"),
        ],
    )
    out = tmp_path / "images_index.xlsx"

    await image_index.build_index(inventory, out, limit=1)
    assert indexed["described"] == 1

    count, _ = await image_index.build_index(inventory, out)

    assert count == 1
    assert indexed["described"] == 2
    assert load_workbook(out)[INDEX_SHEET].max_row == 3


def test_five_images_run_at_once_by_default():
    assert PARALLELISM == 5


async def test_no_more_than_the_parallelism_are_in_flight(tmp_path, monkeypatch):
    """The ceiling holds across save boundaries, not just within one chunk."""
    inventory = _inventory(
        tmp_path,
        [("Grooming", f"{i}.jpg", "зображення", "так", f"id_{i}") for i in range(12)],
    )
    live = {"now": 0, "peak": 0}

    async def slow_download(client, url):
        live["now"] += 1
        live["peak"] = max(live["peak"], live["now"])
        try:
            await asyncio.sleep(0.01)
            buffer = io.BytesIO()
            Image.new("RGB", (8, 8), "red").save(buffer, "JPEG")
            return buffer.getvalue()
        finally:
            live["now"] -= 1

    async def fake_describe(image, model=None):
        return ImageFacts(description="опис")

    monkeypatch.setattr(drive, "download", slow_download)
    monkeypatch.setattr(image_index, "describe", fake_describe)

    count, _ = await image_index.build_index(
        inventory, tmp_path / "out.xlsx", parallelism=3
    )

    assert count == 12
    assert live["peak"] == 3


async def test_a_parallelism_above_the_save_cadence_is_not_throttled(
    tmp_path, monkeypatch
):
    """A chunk must never be smaller than what may run at once."""
    inventory = _inventory(
        tmp_path,
        [("Grooming", f"{i}.jpg", "зображення", "так", f"id_{i}") for i in range(16)],
    )
    live = {"now": 0, "peak": 0}

    async def slow_download(client, url):
        live["now"] += 1
        live["peak"] = max(live["peak"], live["now"])
        try:
            await asyncio.sleep(0.01)
            buffer = io.BytesIO()
            Image.new("RGB", (8, 8), "red").save(buffer, "JPEG")
            return buffer.getvalue()
        finally:
            live["now"] -= 1

    async def fake_describe(image, model=None):
        return ImageFacts(description="опис")

    monkeypatch.setattr(drive, "download", slow_download)
    monkeypatch.setattr(image_index, "describe", fake_describe)

    await image_index.build_index(inventory, tmp_path / "out.xlsx", parallelism=16)

    assert live["peak"] == 16


async def test_a_folder_with_nothing_recommended_says_so(tmp_path, indexed):
    """Six real folders carry no pilot mark; 'already indexed' would be a lie."""
    inventory = _inventory(
        tmp_path,
        [
            ("B2B", "a.jpg", "зображення", None, "id_a"),
            ("B2B", "b.jpg", "зображення", None, "id_b"),
            ("Grooming", "c.jpg", "зображення", "так", "id_c"),
        ],
    )

    with pytest.raises(ValueError, match="none is marked"):
        await image_index.build_index(inventory, tmp_path / "out.xlsx", folder="B2B")

    # ...and --index-all is exactly what the message tells you to reach for.
    count, _ = await image_index.build_index(
        inventory, tmp_path / "out.xlsx", only_recommended=False, folder="B2B"
    )
    assert count == 2


async def test_an_unknown_folder_lists_the_ones_there_are(tmp_path, indexed):
    inventory = _inventory(
        tmp_path, [("Grooming", "a.jpg", "зображення", "так", "id_a")]
    )

    with pytest.raises(ValueError, match="Grooming"):
        await image_index.build_index(inventory, tmp_path / "out.xlsx", folder="Nope")


async def test_an_index_that_is_already_complete_is_not_an_error(tmp_path, indexed):
    """That is the one case the 'nothing left to index' message is for."""
    inventory = _inventory(
        tmp_path, [("Grooming", "a.jpg", "зображення", "так", "id_a")]
    )
    out = tmp_path / "out.xlsx"

    assert await image_index.build_index(inventory, out) == (1, [])
    assert await image_index.build_index(inventory, out) == (0, [])


async def test_a_parallelism_below_one_is_refused(tmp_path):
    inventory = _inventory(
        tmp_path, [("Grooming", "a.jpg", "зображення", "так", "id_a")]
    )
    with pytest.raises(ValueError, match="at least 1"):
        await image_index.build_index(inventory, tmp_path / "out.xlsx", parallelism=0)


async def test_an_image_that_fails_is_reported_and_left_for_next_time(
    tmp_path, indexed, monkeypatch
):
    """No row means a later run retries it, and the failure is never silent."""
    inventory = _inventory(
        tmp_path,
        [
            ("Grooming", "good.jpg", "зображення", "так", "id_a"),
            ("Grooming", "gone.jpg", "зображення", "так", "id_b"),
        ],
    )
    out = tmp_path / "images_index.xlsx"

    async def half_broken(client, url):
        if url.endswith("id_b/view?usp=drivesdk"):
            raise RuntimeError("404 not found")
        buffer = io.BytesIO()
        Image.new("RGB", (64, 64), "red").save(buffer, "JPEG")
        return buffer.getvalue()

    monkeypatch.setattr(drive, "download", half_broken)
    count, failures = await image_index.build_index(inventory, out)

    assert count == 1
    assert failures and "gone.jpg" in failures[0] and "404" in failures[0]
    assert [c.value for c in load_workbook(out)[INDEX_SHEET][2]][1] == "good.jpg"
