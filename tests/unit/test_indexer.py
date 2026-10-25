"""The indexer: reading each spreadsheet, and embedding only what is not stored yet."""

import pytest
from openpyxl import Workbook

from routines.vector_indexer import indexer
from routines.vector_indexer.indexer import (
    IndexDocument,
    index_collection,
    read_content,
    read_instagram_messages,
    read_products,
)

pytestmark = pytest.mark.unit


def _sheet(path, rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(path)
    return ws


def test_products_take_the_first_filled_of_repeated_columns(tmp_path):
    path = tmp_path / "products.xlsx"
    header = ["Артикул", "Название (UA)", "Название модификации (UA)", "Цена", "Бренд", "Раздел",
              "Описание товара (UA)", "Вес", "Об`єм", "Розмір", "Вага", "Розмір", "Об'єм"]  # fmt: skip
    row = ["ND/1", "Бальзам для лап", "Бальзам 30мл", "295.00", "Natural Dog", "Догляд",
           "<p><b>Paw</b> &amp; soother</p>", "0", None, None, "75 г", "M", "30мл"]  # fmt: skip
    _sheet(path, [header, row, [None] * len(header)])

    (doc,) = read_products(path)

    assert doc.id == "ND_1"
    assert doc.data["price"] == 295.0
    assert doc.data["description"] == "Paw & soother"
    assert (doc.data["volume"], doc.data["size"], doc.data["weight"]) == ("30мл", "M", "75 г")
    assert doc.data["name"] == "Бальзам 30мл"
    assert doc.data["embeddingText"].split("\n")[:3] == ["ND/1", "Бальзам 30мл", "Natural Dog"]
    assert "distance" not in doc.data


def test_products_read_the_ukrainian_b2b_export(tmp_path):
    path = tmp_path / "products.xlsx"
    header = ["ID товару/послуги", "Товар/Послуга", "Виробник", "SKU", "Ціна", "Вага", "Опис",
              "Категорія", "Зображення"]  # fmt: skip
    row = ["ND-HS", "Холдер для стіків", "Natural Dog Company", "ND-HS", 120.5, 0.05,
           "<p>Холдер</p>", "Лапи і кігті",
           "https://x.ua/images/1.jpg, https://x.ua/images/2.jpg"]  # fmt: skip
    _sheet(path, [header, row])

    (doc,) = read_products(path)

    assert doc.id == "ND-HS"
    assert (doc.data["name"], doc.data["brand"]) == ("Холдер для стіків", "Natural Dog Company")
    assert (doc.data["price"], doc.data["weight"]) == (120.5, "0.05")
    assert (doc.data["description"], doc.data["category"]) == ("Холдер", "Лапи і кігті")
    # The first of the listed images, and not part of what is embedded.
    assert doc.data["imageUrl"] == "https://x.ua/images/1.jpg"
    assert "x.ua" not in doc.data["embeddingText"]


def test_products_without_a_sku_column_fail_rather_than_index_nothing(tmp_path):
    path = tmp_path / "products.xlsx"
    _sheet(path, [["Назва", "Ціна"], ["Бальзам", 1]])

    with pytest.raises(ValueError, match="no SKU column"):
        read_products(path)


def test_messages_skip_instagram_notices_and_get_stable_ids(tmp_path):
    path = tmp_path / "dm.xlsx"
    header = ["Conversation", "Thread Folder", "Sender", "Timestamp UTC", "Message"]
    rows = [
        header,
        ["c", "t1", "Оля", "2026-02-05 11:33:24.577000", "Цікавить прайс?"],
        ["c", "t1", "Оля", "2026-02-05 11:34:00", "Liked a message"],
        ["c", "t1", "Оля", "2026-02-05 11:35:00", "[Фото]"],
        ["c", "t1", "Оля", "2026-02-05 11:36:00", None],
    ]
    _sheet(path, rows)

    (doc,) = read_instagram_messages(path)

    assert doc.data["text"] == doc.data["embeddingText"] == "Цікавить прайс?"
    assert doc.data["timestamp"].tzinfo is not None
    assert read_instagram_messages(path)[0].id == doc.id


def test_messages_are_joined_per_conversation_sender_and_kyiv_day(tmp_path):
    path = tmp_path / "dm.xlsx"
    shop = "HealthyDoggo"
    rows = [
        ["Conversation", "Thread Folder", "Sender", "Timestamp UTC", "Message"],
        ["c", "t1", "Оля", "2026-02-05 10:05:00", "а на велику породу?"],
        ["c", "t1", "Оля", "2026-02-05 09:00:00", "Цікавить шампунь"],
        ["c", "t1", shop, "2026-02-05 09:30:00", "Добрий день"],
        ["c", "t2", shop, "2026-02-05 09:40:00", "Вітаємо"],
        # 22:30 UTC on the 5th is 00:30 on the 6th in Kyiv: a new day.
        ["c", "t1", "Оля", "2026-02-05 22:30:00", "дякую"],
    ]
    _sheet(path, rows)

    docs = {(d.data["sender"], d.data["text"]): d for d in read_instagram_messages(path)}

    assert set(docs) == {
        ("Оля", "Цікавить шампунь\nа на велику породу?"),
        (shop, "Добрий день"),
        (shop, "Вітаємо"),
        ("Оля", "дякую"),
    }
    joined = docs[("Оля", "Цікавить шампунь\nа на велику породу?")].data
    assert joined["timestamp"].isoformat() == "2026-02-05T09:00:00+00:00"
    assert joined["embeddingText"] == joined["text"]
    assert len({d.id for d in docs.values()}) == 4


def test_content_reads_links_from_hyperlinks_and_blanks_not_applicable(tmp_path):
    path = tmp_path / "images_index.xlsx"
    header = ["#", "Файл", "Папка", "Google Drive", "Опис", "Screenshots", "AudioTranscribe",
              "Продукт", "Бренд", "Порода", "Теги", "Тип", "Текст на фото", "Розмір, KB"]  # fmt: skip
    row = [1, "dog.mp4", "UGC", "Відкрити", "Шпіц біжить", "Відкрити", "Привіт",
           "—", "—", "шпіц", "собака, шпіц", "лайфстайл", "—", 108.1]  # fmt: skip
    ws = _sheet(path, [header, row])
    ws["D2"].hyperlink = "https://drive.google.com/file/d/abc_123-X/view?usp=drivesdk"
    ws["F2"].hyperlink = "https://drive.google.com/drive/folders/shots"
    ws.parent.save(path)

    (doc,) = read_content(path)

    assert doc.id == "abc_123-X"
    assert doc.data["driveUrl"].startswith("https://drive.google.com/file/d/abc_123-X")
    assert doc.data["screenshotsPath"] == "https://drive.google.com/drive/folders/shots"
    assert doc.data["product"] == doc.data["brand"] == doc.data["textOnImage"] == ""
    assert doc.data["sizeKb"] == 108.1
    assert doc.data["hasDog"] is True
    assert doc.data["embeddingText"] == "Шпіц біжить\nПривіт\nшпіц\nсобака, шпіц\nлайфстайл"


class _Firestore:
    def __init__(self, stored=()):
        self.stored = set(stored)
        self.writes = []

    async def existing_ids(self, collection):
        return set(self.stored)

    async def upsert(self, collection, documents):
        self.writes.extend(documents)


class _Embedder:
    def __init__(self):
        self.texts = []

    async def embed_documents(self, texts):
        self.texts.extend(texts)
        return [[0.1] for _ in texts]


def _doc(doc_id, text="t"):
    return IndexDocument(doc_id, {"embeddingText": text})


async def test_index_skips_stored_and_textless_rows_and_writes_vectors():
    firestore, embedder = _Firestore(stored={"a"}), _Embedder()
    docs = [_doc("a"), _doc("b"), _doc("c", ""), _doc("d"), _doc("d")]

    summary = await index_collection("products", docs, firestore, embedder)

    assert (summary.indexed, summary.already_stored, summary.no_text) == (2, 1, 1)
    assert [doc_id for doc_id, _ in firestore.writes] == ["b", "d"]
    assert firestore.writes[0][1] == {"embeddingText": "t", "embeddingVector": [0.1]}


def _record_sleeps(monkeypatch):
    sleeps = []

    async def sleep(seconds):
        sleeps.append(seconds)

    monkeypatch.setattr(indexer.asyncio, "sleep", sleep)
    return sleeps


async def test_force_reindexes_and_limit_caps_the_run(monkeypatch):
    monkeypatch.setattr(indexer, "CHUNK", 1)
    _record_sleeps(monkeypatch)
    firestore, embedder = _Firestore(stored={"a"}), _Embedder()

    summary = await index_collection(
        "products", [_doc("a"), _doc("b"), _doc("c")], firestore, embedder, limit=2, force=True
    )

    assert summary.indexed == 2
    assert [doc_id for doc_id, _ in firestore.writes] == ["a", "b"]


@pytest.mark.parametrize(("docs", "expected"), [(3, 2), (1, 0), (0, 0)])
async def test_the_indexer_pauses_between_chunks_but_not_after_the_last(
    monkeypatch, docs, expected
):
    monkeypatch.setattr(indexer, "CHUNK", 1)
    sleeps = _record_sleeps(monkeypatch)

    await index_collection(
        "products", [_doc(str(i)) for i in range(docs)], _Firestore(), _Embedder()
    )

    assert sleeps == [indexer.CHUNK_DELAY] * expected


@pytest.mark.parametrize(
    ("breed", "tags", "expected"),
    [
        ("шпіц", "", True),
        ("шпіц", "без тварини", True),
        ("", "Собака, лапа, біла шерсть", True),
        ("", "засіб, догляд за собаками", False),
        ("", "собака, догляд за собакою, без тварини", False),
        ("", "собака відсутня, пляшка", False),
        ("", "кіт, диван", False),
        ("", "", False),
    ],
)
def test_a_dog_comes_from_the_breed_or_an_exact_dog_tag(breed, tags, expected):
    assert indexer._has_dog(breed, tags) is expected
